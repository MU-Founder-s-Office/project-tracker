#!/usr/bin/env python3
import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
from openpyxl.utils.datetime import from_excel


PROJECT_ID_RE = re.compile(r"^P-\d{2,}$")
URL_RE = re.compile(r"(https?://[^\s|,]+|www\.[^\s|,]+)", re.IGNORECASE)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return " ".join(value.replace("\n", " ").split())
    return value


def excel_date(value, workbook):
    if value in ("", None):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value, workbook.epoch).date().isoformat()
        except Exception:
            return None
    return str(value)


def percent_value(value):
    if value in ("", None):
        return 0
    if isinstance(value, (int, float)):
        return round(value * 100) if value <= 1 else round(value)
    text = str(value).strip()
    match = re.search(r"-?\d+(\.\d+)?", text)
    if not match:
        return 0
    return round(float(match.group(0)))


def duration_days(start_date, end_date):
    if not start_date or not end_date:
        return 0
    try:
        start = datetime.fromisoformat(start_date)
        end = datetime.fromisoformat(end_date)
        return max((end - start).days, 0)
    except Exception:
        return 0


def normalize_url(value):
    if not value:
        return ""
    text = str(value).strip()
    if text.lower().startswith("www."):
        return f"https://{text}"
    return text


def link_label(url):
    try:
        parsed = urlparse(url)
        host = parsed.netloc or parsed.path
        if "docs.google.com" in host:
            return "Google Doc"
        if "netlify.app" in host:
            return "Netlify App"
        if "vercel.app" in host:
            return "Vercel App"
        if "claude.ai" in host:
            return "Claude Artifact"
        return host.replace("www.", "") or "Project Link"
    except Exception:
        return "Project Link"


def extract_links(cell):
    links = []
    text = clean(cell.value)

    if cell.hyperlink and cell.hyperlink.target:
        target = normalize_url(cell.hyperlink.target)
        if target:
            links.append({"label": link_label(target), "url": target, "source": "hyperlink"})

    if isinstance(text, str):
        for match in URL_RE.finditer(text):
            url = normalize_url(match.group(0).rstrip(").]"))
            if url and all(item["url"] != url for item in links):
                links.append({"label": link_label(url), "url": url, "source": "text"})

    return links


def extract_links_from_cells(cells):
    links = []
    for cell in cells:
        for link in extract_links(cell):
            if link["url"] and all(item["url"] != link["url"] for item in links):
                links.append(link)
    return links


def merge_unique(existing, incoming):
    merged = list(existing or [])
    for item in incoming or []:
        if isinstance(item, dict):
            key = item.get("url") or json.dumps(item, sort_keys=True)
            if all((entry.get("url") if isinstance(entry, dict) else entry) != key for entry in merged):
                merged.append(item)
        elif item and item not in merged:
            merged.append(item)
    return merged


def title_case_workstream(value):
    text = clean(value)
    if not isinstance(text, str) or not text:
        return ""
    return text.title().replace("Nps", "NPS").replace("Ai", "AI")


def split_note_and_refs(value):
    text = clean(value)
    if not isinstance(text, str) or "|" not in text:
        return text, []
    parts = [part.strip() for part in text.split("|") if part.strip()]
    note = parts[0] if parts else ""
    refs = [part for part in parts[1:] if not URL_RE.search(part)]
    return note, refs


def read_projects_from_breakdown(workbook):
    sheet = workbook["Breakdown"]
    projects = []
    current_workstream = ""

    for row in sheet.iter_rows(min_row=5, max_col=sheet.max_column):
        section_value = clean(row[1].value) if len(row) > 1 else ""
        project_id = clean(row[3].value) if len(row) > 3 else ""

        if section_value and not project_id and str(section_value).strip().isupper():
            current_workstream = title_case_workstream(section_value)
            continue

        if not PROJECT_ID_RE.match(str(project_id).strip()):
            continue

        project_name = clean(row[2].value) if len(row) > 2 else ""
        status = clean(row[4].value) if len(row) > 4 else ""
        priority = clean(row[5].value) if len(row) > 5 else ""
        progress = clean(row[6].value) if len(row) > 6 else ""
        top_risk = clean(row[7].value) if len(row) > 7 else ""
        note = clean(row[8].value) if len(row) > 8 else ""
        start_date = excel_date(clean(row[9].value) if len(row) > 9 else "", workbook)
        end_date = excel_date(clean(row[10].value) if len(row) > 10 else "", workbook)
        link_cell = row[11] if len(row) > 11 else None
        link_text = clean(link_cell.value) if link_cell else ""
        links = extract_links_from_cells(row)
        references = []
        if isinstance(link_text, str) and link_text and not URL_RE.search(link_text):
            references.append(link_text)

        extra_fields = {}

        projects.append(
            {
                "id": str(project_id).strip(),
                "project": str(project_name).strip(),
                "workstream": current_workstream,
                "status": status,
                "priority": priority,
                "startDate": start_date,
                "endDate": end_date,
                "progress": percent_value(progress),
                "topRisk": top_risk,
                "notes": note,
                "references": references,
                "links": links,
                "rag": "",
                "durationDays": duration_days(start_date, end_date),
                "extraFields": extra_fields,
            }
        )

    return projects


def read_projects(workbook):
    sheet = workbook["Projects"]
    headers = [clean(cell.value) for cell in sheet[4]]
    projects = []

    for row in sheet.iter_rows(min_row=5):
        row_values = {headers[idx]: clean(cell.value) for idx, cell in enumerate(row) if idx < len(headers) and headers[idx]}
        project_id = str(row_values.get("ID", "")).strip()
        if not PROJECT_ID_RE.match(project_id):
            continue

        notes_cell = row[9]
        note, references = split_note_and_refs(row_values.get("Notes / Links", ""))
        links = extract_links_from_cells(row)

        extra_fields = {}
        known = {
            "ID",
            "Project",
            "Workstream",
            "Status",
            "Priority",
            "Start Date",
            "End Date",
            "% Done",
            "Top Risk",
            "Notes / Links",
            "RAG",
            "Dur (d)",
        }
        for key, value in row_values.items():
            if key not in known and value not in ("", None):
                extra_fields[key] = value

        projects.append(
            {
                "id": project_id,
                "project": row_values.get("Project", ""),
                "workstream": row_values.get("Workstream", ""),
                "status": row_values.get("Status", ""),
                "priority": row_values.get("Priority", ""),
                "startDate": excel_date(row_values.get("Start Date"), workbook),
                "endDate": excel_date(row_values.get("End Date"), workbook),
                "progress": percent_value(row_values.get("% Done")),
                "topRisk": row_values.get("Top Risk", ""),
                "notes": note,
                "references": references,
                "links": links,
                "rag": row_values.get("RAG", ""),
                "durationDays": row_values.get("Dur (d)", 0) or duration_days(
                    excel_date(row_values.get("Start Date"), workbook),
                    excel_date(row_values.get("End Date"), workbook),
                ),
                "extraFields": extra_fields,
            }
        )

    if "Breakdown" in workbook.sheetnames:
        breakdown_projects = {project["id"]: project for project in read_projects_from_breakdown(workbook)}
        for project in projects:
            fallback = breakdown_projects.get(project["id"])
            if not fallback:
                continue

            for key in [
                "project",
                "workstream",
                "status",
                "priority",
                "startDate",
                "endDate",
                "progress",
                "topRisk",
                "notes",
                "rag",
                "durationDays",
            ]:
                if project.get(key) in ("", None, 0) and fallback.get(key) not in ("", None, 0):
                    project[key] = fallback[key]

            project["links"] = merge_unique(project.get("links", []), fallback.get("links", []))
            project["references"] = merge_unique(project.get("references", []), fallback.get("references", []))
            project["extraFields"] = {**fallback.get("extraFields", {}), **project.get("extraFields", {})}

    return projects


def read_risks(workbook):
    if "Risks" not in workbook.sheetnames:
        return []

    sheet = workbook["Risks"]
    headers = [clean(cell.value) for cell in sheet[3]]
    risks = []
    for row in sheet.iter_rows(min_row=4):
        values = {headers[idx]: clean(cell.value) for idx, cell in enumerate(row) if idx < len(headers) and headers[idx]}
        project_id = str(values.get("ID", "")).strip()
        if not PROJECT_ID_RE.match(project_id):
            continue
        risks.append(
            {
                "id": project_id,
                "project": values.get("Project", ""),
                "workstream": values.get("Workstream", ""),
                "topRisk": values.get("Top Risk", ""),
                "likelihood": values.get("Likelihood", values.get("L", 0)) or 0,
                "impact": values.get("Impact", values.get("I", 0)) or 0,
                "severity": values.get("Severity", values.get("Sev", 0)) or 0,
                "mitigation": values.get("Mitigation", ""),
                "owner": values.get("Owner", values.get("Workstream", "")),
            }
        )
    return risks


def build_dataset(source):
    workbook = openpyxl.load_workbook(source, data_only=True)
    projects = read_projects(workbook)
    risks = read_risks(workbook)
    risk_by_id = {risk["id"]: risk for risk in risks}

    for project in projects:
        risk = risk_by_id.get(project["id"], {})
        if risk:
            if not project.get("project"):
                project["project"] = risk.get("project", "")
            if not project.get("workstream"):
                project["workstream"] = risk.get("workstream", "")
            if not project.get("topRisk"):
                project["topRisk"] = risk.get("topRisk", "")
            if not project.get("notes"):
                project["notes"] = risk.get("mitigation", "")
        project["risk"] = risk

    workstreams = sorted({project["workstream"] for project in projects if project["workstream"]})
    statuses = sorted({project["status"] for project in projects if project["status"]})
    priorities = sorted({project["priority"] for project in projects if project["priority"]})

    return {
        "metadata": {
            "sourceFile": str(Path(source).name),
            "generatedAt": datetime.now(timezone.utc).isoformat(),
            "projectCount": len(projects),
            "workstreams": workstreams,
            "statuses": statuses,
            "priorities": priorities,
        },
        "projects": projects,
        "risks": risks,
    }


def main():
    parser = argparse.ArgumentParser(description="Import the MU project tracker workbook into dashboard JSON.")
    parser.add_argument("source", help="Path to MU_Project_Tracker workbook")
    parser.add_argument(
        "--output",
        default=str(Path(__file__).resolve().parents[1] / "data" / "projects.json"),
        help="Output JSON path",
    )
    args = parser.parse_args()

    dataset = build_dataset(args.source)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(dataset, indent=2), encoding="utf-8")

    print(f"Imported {dataset['metadata']['projectCount']} projects from {args.source}")
    print(f"Wrote {output}")


if __name__ == "__main__":
    main()
